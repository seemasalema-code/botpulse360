#!/usr/bin/env python3
"""Build, reconcile, partition and archive the dashboard snapshot."""
from __future__ import annotations
import argparse, datetime as dt, gzip, hashlib, json, os, re
from pathlib import Path
from openpyxl import load_workbook

RULES = {
    "Chatbot Projects": ({"Client", "Chatbot"}, 10000),
    "Chatbot R&M": ({"Chatbot", "Month"}, 30000),
    "WA_Consumables": ({"Client", "Month", "Revenue", "Cost"}, 30000),
    "RCS_Consumables": ({"Client", "Month", "Revenue", "Cost"}, 30000),
}
MONTH_TABS = ("Chatbot R&M", "WA_Consumables", "RCS_Consumables")
REVENUE_HEADERS = {
    "Chatbot R&M": ("Total Netcore Revenue", "Netcore Total Revenue", "Total Revenue"),
    "WA_Consumables": ("Revenue",), "RCS_Consumables": ("Revenue",),
}

def serialise(value):
    if value is None: return ""
    if isinstance(value, (dt.datetime, dt.date)): return value.strftime("%d/%m/%Y")
    if isinstance(value, float) and value.is_integer(): return int(value)
    return value

def trim_sheet(ws, limit):
    rows, empty = [], 0
    for row_number, row in enumerate(ws.iter_rows(values_only=True), 1):
        if row_number > limit: raise ValueError(f"{ws.title} exceeds the {limit:,}-row safety limit")
        values = [serialise(v) for v in row]
        while values and values[-1] == "": values.pop()
        if not values:
            empty += 1
            if empty >= 100: break
            rows.append([]); continue
        empty = 0; rows.append(values)
    while rows and not rows[-1]: rows.pop()
    return rows

def header_info(tab, rows):
    required = RULES[tab][0]
    for index, row in enumerate(rows[:10]):
        mapping = {str(v).strip(): pos for pos, v in enumerate(row) if v not in (None, "")}
        if required.issubset(mapping): return index, mapping
    raise ValueError(f"{tab} is missing required headers: {', '.join(sorted(required))}")

def number(value):
    if isinstance(value, (int, float)): return float(value)
    cleaned = re.sub(r"[^0-9.()-]", "", str(value or "")).replace("(", "-").replace(")", "")
    try: return float(cleaned) if cleaned else 0.0
    except ValueError: return 0.0

def parse_month(value):
    text = str(value or "").strip()
    for fmt in ("%d/%m/%Y", "%b-%y", "%b %Y", "%Y-%m-%d"):
        try: return dt.datetime.strptime(text, fmt).date().replace(day=1)
        except ValueError: pass
    return None

def fy(date):
    start = date.year if date.month >= 4 else date.year - 1
    return f"FY{start}-{str(start + 1)[-2:]}"

def read_previous(path):
    if not path.exists(): return {}
    text = path.read_text(encoding="utf-8"); prefix = "window.EMBEDDED_SHEETS_CURRENT = "
    try: return json.loads(text[len(prefix):].rsplit(";", 1)[0]) if text.startswith(prefix) else {}
    except (ValueError, json.JSONDecodeError): return {}

def metrics(tab, rows):
    header_row, headers = header_info(tab, rows)
    records = [r for r in rows[header_row + 1:] if any(v not in (None, "") for v in r)]
    revenue_col = next((headers[h] for h in REVENUE_HEADERS.get(tab, ()) if h in headers), None)
    revenue = sum(number(r[revenue_col]) for r in records if revenue_col is not None and len(r) > revenue_col)
    return {"rows": len(records), "revenue": round(revenue, 2)}

def reconcile(snapshot, previous):
    current, prior, comparison, warnings, failures = {}, {}, {}, [], []
    for tab, (_, limit) in RULES.items():
        current[tab] = metrics(tab, snapshot[tab])
        if previous.get(tab): prior[tab] = metrics(tab, previous[tab])
        if current[tab]["rows"] >= limit * .8:
            warnings.append(f"{tab} is at {current[tab]['rows'] / limit:.0%} of its safety limit")
        old = prior.get(tab, {"rows": 0, "revenue": 0})
        comparison[tab] = {"rowDelta": current[tab]["rows"] - old["rows"], "revenueDelta": round(current[tab]["revenue"] - old["revenue"], 2)}
        if old["rows"] and current[tab]["rows"] < old["rows"] * .6:
            failures.append(f"{tab} rows fell over 40% ({old['rows']} to {current[tab]['rows']})")
        if old["revenue"] > 0 and current[tab]["revenue"] < old["revenue"] * .6:
            failures.append(f"{tab} revenue fell over 40% ({old['revenue']:.2f} to {current[tab]['revenue']:.2f})")
    return current, prior, comparison, warnings, failures

def write_partitions(snapshot, root):
    partitions, archives = {}, {}
    for tab in MONTH_TABS:
        hi, headers = header_info(tab, snapshot[tab]); header = snapshot[tab][hi]; mi = headers["Month"]
        for row in snapshot[tab][hi + 1:]:
            date = parse_month(row[mi] if len(row) > mi else "")
            if not date: continue
            key = (tab, fy(date), date.strftime("%Y-%m"))
            partitions.setdefault(key, [header]).append(row)
            archives.setdefault(fy(date), {}).setdefault(tab, [header]).append(row)
    for (tab, year, month), rows in partitions.items():
        path = root / "partitions" / tab.replace(" ", "_").replace("&", "and") / year / f"{month}.json"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(rows, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    for year, tabs in archives.items():
        path = root / "archive" / f"{year}.json.gz"; path.parent.mkdir(parents=True, exist_ok=True)
        with gzip.open(path, "wt", encoding="utf-8") as stream: json.dump(tabs, stream, ensure_ascii=False, separators=(",", ":"))

def main():
    parser = argparse.ArgumentParser(); parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, default=Path("data.js")); parser.add_argument("--report", type=Path, default=Path("snapshot-report.json")); parser.add_argument("--data-dir", type=Path, default=Path("data")); args = parser.parse_args()
    previous = read_previous(args.output); workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    missing = [tab for tab in RULES if tab not in workbook.sheetnames]
    if missing: raise ValueError("Workbook is missing required tabs: " + ", ".join(missing))
    snapshot = {}
    for tab, (_, limit) in RULES.items(): snapshot[tab] = trim_sheet(workbook[tab], limit); header_info(tab, snapshot[tab])
    current, prior, comparison, warnings, failures = reconcile(snapshot, previous)
    report = {"generatedAt": dt.datetime.now(dt.timezone.utc).isoformat(), "sourceFile": args.workbook.name, "current": current, "previous": prior, "comparison": comparison, "warnings": warnings, "failures": failures}
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    if failures and os.environ.get("ALLOW_SNAPSHOT_ANOMALY") != "1": raise ValueError("Reconciliation failed: " + "; ".join(failures))
    raw = args.workbook.read_bytes(); snapshot["__meta"] = {"sourceFile": args.workbook.name, "sourceSha256": hashlib.sha256(raw).hexdigest(), "refreshedAt": dt.datetime.now(dt.timezone.utc).isoformat(), "mode": "validated-workbook-snapshot", "reconciliation": {"warnings": warnings, "comparison": comparison}}
    write_partitions(snapshot, args.data_dir)
    payload = "window.EMBEDDED_SHEETS_CURRENT = " + json.dumps(snapshot, ensure_ascii=False, separators=(",", ":")) + ";\n"
    temporary = args.output.with_suffix(args.output.suffix + ".tmp"); temporary.write_text(payload, encoding="utf-8"); temporary.replace(args.output)
    print(json.dumps(report, indent=2))

if __name__ == "__main__": main()
